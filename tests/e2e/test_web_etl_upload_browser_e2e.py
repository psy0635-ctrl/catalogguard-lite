"""Chromium coverage for the Streamlit Web ETL CSV/XLSX upload flow."""

import os
import re
import hashlib
import json
import csv
from pathlib import Path

import pytest
from openpyxl import Workbook


pytestmark = pytest.mark.e2e


ROOT = Path(__file__).resolve().parents[2]
FIXTURE_PATH = ROOT / "tests" / "fixtures" / "e2e" / "web_etl_upload_vendor.csv"
STREAMLIT_URL = os.environ.get("E2E_STREAMLIT_URL", "http://127.0.0.1:8501")
SOURCE_FILENAME = os.environ.get("E2E_WEB_UPLOAD_SOURCE_FILENAME", FIXTURE_PATH.name)
E2E_OPERATOR_USERNAME = os.environ.get("E2E_OPERATOR_USERNAME", "")
E2E_OPERATOR_PASSWORD = os.environ.get("E2E_OPERATOR_PASSWORD", "")
API_URL = os.environ.get(
    "E2E_API_URL",
    os.environ.get("CATALOGGUARD_API_BASE_URL", "http://127.0.0.1:8000"),
)
PROFILE_ID = "sample_marketplace_vendor_v1"
PRODUCT_IDS = (
    "CG-E2E-WEB-UPLOAD-BLK-M",
    "CG-E2E-WEB-UPLOAD-WHT-L",
)


def _login_as_operator(page) -> None:
    from playwright.sync_api import expect

    assert E2E_OPERATOR_USERNAME and E2E_OPERATOR_PASSWORD, (
        "E2E_OPERATOR_USERNAME/E2E_OPERATOR_PASSWORD 환경변수가 필요합니다."
    )
    expect(page.locator("body")).to_contain_text(
        "좌측 사이드바에서 로그인한 뒤 이용해 주세요."
    )
    page.get_by_label("아이디").fill(E2E_OPERATOR_USERNAME)
    page.get_by_label("비밀번호").fill(E2E_OPERATOR_PASSWORD)
    page.get_by_role("button", name="로그인", exact=True).click()
    expect(page.get_by_role("tab", name="ETL 적재 이력")).to_be_visible()


def _preserve_browser_failure_artifacts(page) -> None:
    artifact_dir = os.environ.get("E2E_ARTIFACT_DIR", "").strip()
    if not artifact_dir:
        return
    output_dir = Path(artifact_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        page.screenshot(path=str(output_dir / "failure.png"), full_page=True)
        (output_dir / "page.html").write_text(page.content(), encoding="utf-8")
    except Exception:
        pass


def _expected_profile_definition_sha256() -> str:
    from etl.profile_fingerprint import compute_profile_definition_sha256
    from etl.profile_loader import get_profile_path, load_profile

    return compute_profile_definition_sha256(load_profile(get_profile_path(PROFILE_ID)))


def _expected_profile_definition_snapshot() -> dict[str, object]:
    from etl.profile_fingerprint import build_profile_semantic_payload
    from etl.profile_loader import get_profile_path, load_profile

    return build_profile_semantic_payload(load_profile(get_profile_path(PROFILE_ID)))


def _expected_application_commit_sha() -> str | None:
    from etl.application_lineage import resolve_application_commit_sha

    return resolve_application_commit_sha()


def _assert_web_upload_persisted(
    etl_load_run_id: int,
    *,
    source_filename: str,
) -> tuple[str, dict[str, object], str | None]:
    from sqlalchemy import func, select

    from db.models import CatalogProductStaging, ETLLoadRun
    from db.session import create_session_factory

    session_factory = create_session_factory(database_url=os.environ["DATABASE_URL"])
    with session_factory() as session:
        load_run = session.get(ETLLoadRun, etl_load_run_id)
        assert load_run is not None
        assert load_run.source_filename == source_filename
        assert load_run.profile_name == "sample_marketplace_vendor"
        assert load_run.profile_version == "2"
        expected_fingerprint = _expected_profile_definition_sha256()
        expected_snapshot = _expected_profile_definition_snapshot()
        assert load_run.profile_definition_sha256 == expected_fingerprint
        assert load_run.profile_definition_snapshot == expected_snapshot
        assert hashlib.sha256(
            json.dumps(
                expected_snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ).encode("utf-8")
        ).hexdigest() == expected_fingerprint
        expected_application_commit_sha = _expected_application_commit_sha()
        assert load_run.application_commit_sha == expected_application_commit_sha
        assert load_run.initial_source_type == "upload"
        assert load_run.initial_source_ref == source_filename

        product_count = session.scalar(
            select(func.count())
            .select_from(CatalogProductStaging)
            .where(CatalogProductStaging.etl_load_run_id == etl_load_run_id)
        )
        assert product_count == 2
        staging_product_ids = set(
            session.scalars(
                select(CatalogProductStaging.product_id).where(
                    CatalogProductStaging.etl_load_run_id == etl_load_run_id
                )
            )
        )
        assert staging_product_ids == set(PRODUCT_IDS)
    return expected_fingerprint, expected_snapshot, expected_application_commit_sha


def _api_headers(page) -> dict[str, str]:
    response = page.request.post(
        f"{API_URL}/api/v1/auth/login",
        data={"username": E2E_OPERATOR_USERNAME, "password": E2E_OPERATOR_PASSWORD},
    )
    assert response.ok, response.text()
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _run_web_etl_upload_scenario(
    page,
    *,
    fixture_path: Path = FIXTURE_PATH,
    source_filename: str = SOURCE_FILENAME,
) -> None:
    from playwright.sync_api import expect

    assert fixture_path.is_file(), f"Browser E2E fixture is missing: {fixture_path}"

    console_errors: list[str] = []
    page_errors: list[str] = []
    page.on(
        "console",
        lambda message: console_errors.append(message.text)
        if message.type == "error"
        else None,
    )
    page.on("pageerror", lambda error: page_errors.append(str(error)))

    page.goto(STREAMLIT_URL, wait_until="domcontentloaded")
    expect(page.locator("body")).to_contain_text("CatalogGuard Lite")
    expect(page.locator("body")).not_to_contain_text("StreamlitAPIException")

    _login_as_operator(page)
    page.get_by_role("tab", name="ETL 적재 이력").click()
    expect(page.locator("body")).to_contain_text("ETL 실행")

    profile_selector = page.get_by_role("combobox", name="ETL 실행 프로필")
    expect(profile_selector).to_be_visible()
    profile_selector.click()
    page.get_by_role("option", name="마켓플레이스 공급사 샘플", exact=True).click()
    expect(page.locator("body")).to_contain_text("sample_marketplace_vendor")
    expect(page.locator("body")).to_contain_text("프로필 버전")
    expect(page.locator("body")).to_contain_text(re.compile(r"프로필 버전\s*2"))

    uploader_dropzone = page.get_by_label("공급사 CSV / XLSX 파일", exact=True)
    file_input = uploader_dropzone.locator('input[type="file"]')
    expect(file_input).to_have_count(1)
    file_input.set_input_files(str(fixture_path))
    run_button = page.get_by_role("button", name="ETL 실행", exact=True)
    expect(run_button).to_be_enabled()
    run_button.click()

    success_pattern = re.compile(r"ETL 실행이 완료되었습니다\. 적재 배치 ID:\s*(\d+)")
    expect(page.locator("body")).to_contain_text(success_pattern)
    expect(page.locator("body")).not_to_contain_text(
        "동일한 입력으로 이미 처리된 적재 배치입니다."
    )
    body_text = page.locator("body").inner_text()
    match = success_pattern.search(body_text)
    assert match is not None
    etl_load_run_id = int(match.group(1))
    expected_fingerprint, expected_snapshot, expected_application_commit_sha = _assert_web_upload_persisted(
        etl_load_run_id,
        source_filename=source_filename,
    )
    detail_response = page.request.get(
        f"{API_URL}/api/v1/etl-loads/{etl_load_run_id}", headers=_api_headers(page)
    )
    assert detail_response.ok, detail_response.text()
    assert detail_response.json()["profile_definition_sha256"] == expected_fingerprint
    assert detail_response.json()["profile_definition_snapshot"] == expected_snapshot
    assert detail_response.json()["application_commit_sha"] == expected_application_commit_sha

    expect(page.locator("body")).to_contain_text("ETL 적재 이력")
    page.get_by_label("원본 파일명").fill(source_filename)
    page.get_by_label("공급사 프로필").fill("sample_marketplace_vendor")
    page.get_by_role("button", name="조회", exact=True).click()
    expect(page.locator("body")).to_contain_text(source_filename)
    expect(page.locator("body")).to_contain_text("sample_marketplace_vendor")

    batch_selector = page.get_by_role("combobox", name="적재 배치 선택")
    batch_selector.click()
    page.get_by_role(
        "option", name=re.compile(rf"{etl_load_run_id} · {re.escape(source_filename)}")
    ).click()
    page.get_by_role("button", name="상세 조회", exact=True).click()

    expect(page.locator("body")).to_contain_text("적재 배치 상세")
    expect(page.locator("body")).to_contain_text(f"원본 파일명: {source_filename}")
    expect(page.locator("body")).to_contain_text(
        "공급사 프로필: sample_marketplace_vendor"
    )
    expect(page.locator("body")).to_contain_text("프로필 버전: 2")
    expect(page.locator("body")).to_contain_text("프로필 정의 SHA-256")
    expect(page.locator("body")).to_contain_text(expected_fingerprint)
    expect(page.locator("body")).to_contain_text("실행 당시 저장된 Profile 정의")
    expect(page.locator("body")).to_contain_text("sku_code")
    expect(page.locator("body")).to_contain_text("애플리케이션 Commit SHA")
    expect(page.locator("body")).to_contain_text(
        expected_application_commit_sha or "알 수 없음 (legacy/미확인)"
    )
    expect(page.locator("body")).to_contain_text("정상 적재")
    expect(page.locator("body")).to_contain_text("전체 입력")
    expect(page.locator("body")).to_contain_text(PRODUCT_IDS[0])

    assert not console_errors, f"Unexpected browser console errors: {console_errors}"
    assert not page_errors, f"Unexpected browser page errors: {page_errors}"


def test_web_etl_upload_flow_in_real_browser(page):
    try:
        _run_web_etl_upload_scenario(page)
    except BaseException:
        _preserve_browser_failure_artifacts(page)
        raise


def _snapshot_database(factory):
    # Import for its registration side effect, even when this test runs alone.
    import db.models
    from db.base import Base
    from sqlalchemy import select

    tables = Base.metadata.sorted_tables
    required_tables = {
        "users", "etl_load_runs", "etl_rejected_rows", "catalog_products_staging"
    }
    registered_tables = {table.name for table in tables}
    missing_tables = required_tables - registered_tables
    assert not missing_tables, (
        f"Database snapshot is missing required ORM tables: {sorted(missing_tables)}"
    )
    with factory() as session:
        return {
            table.name: sorted(repr(tuple(row)) for row in session.execute(select(table)))
            for table in tables
        }


def test_web_etl_preflight_in_real_browser_preserves_database(page):
    from playwright.sync_api import expect
    from db.session import create_session_factory

    factory = create_session_factory(database_url=os.environ["DATABASE_URL"])
    def snapshot():
        return _snapshot_database(factory)
    try:
        page.goto(STREAMLIT_URL, wait_until="domcontentloaded")
        _login_as_operator(page)
        page.get_by_role("tab", name="ETL 적재 이력").click()
        page.get_by_role("combobox", name="ETL 실행 프로필").click()
        page.get_by_role("option", name="마켓플레이스 공급사 샘플", exact=True).click()
        button = page.get_by_role("button", name="사전 검증", exact=True)
        expect(button).to_be_disabled()
        page.get_by_label("공급사 CSV / XLSX 파일", exact=True).locator('input[type="file"]').set_input_files(str(FIXTURE_PATH))
        expect(button).to_be_enabled()
        before = snapshot()
        button.click()
        expect(page.locator("body")).to_contain_text("사전 검증 완료")
        expect(page.locator("body")).to_contain_text("정상 변환 예상")
        expect(page.locator("body")).to_contain_text("Reject 예상 0건")
        expect(page.locator("body")).to_contain_text("아직 데이터베이스에 저장하지 않았습니다.")
        expect(page.get_by_role("button", name="ETL 실행", exact=True)).to_be_enabled()
        assert snapshot() == before
    except BaseException:
        _preserve_browser_failure_artifacts(page)
        raise
    finally:
        factory.kw["bind"].dispose()


def test_web_etl_preflight_recovery_after_correcting_csv_in_real_browser(page, tmp_path):
    from playwright.sync_api import expect
    from db.session import create_session_factory

    source = ROOT / "tests" / "fixtures" / "etl" / "sample_marketplace_vendor_mixed.csv"
    original_bytes = source.read_bytes()
    before_file = tmp_path / "before" / "supplier.csv"
    after_file = tmp_path / "after" / "supplier.csv"
    before_file.parent.mkdir()
    after_file.parent.mkdir()
    before_file.write_bytes(original_bytes)
    with before_file.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames
        original_rows = list(reader)
    corrected_rows = [dict(row) for row in original_rows]
    rejected_rows = [row for row in corrected_rows if row["sku_code"] == "SKU-100-RED-S"]
    assert len(rejected_rows) == 1
    rejected_row = rejected_rows[0]
    assert (rejected_row["regular_price"], rejected_row["available_qty"]) == ("가격문의", "-1")
    rejected_row.update(regular_price="59000", available_qty="3")
    with after_file.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(corrected_rows)
    with after_file.open(encoding="utf-8", newline="") as handle:
        assert list(csv.DictReader(handle)) == corrected_rows
    changed_cells = [
        (row["sku_code"], column)
        for row, corrected in zip(original_rows, corrected_rows, strict=True)
        for column in fieldnames
        if row[column] != corrected[column]
    ]
    assert changed_cells == [
        ("SKU-100-RED-S", "regular_price"), ("SKU-100-RED-S", "available_qty")
    ]
    assert before_file.name == after_file.name
    assert before_file.read_bytes() != after_file.read_bytes()
    assert hashlib.sha256(original_bytes).digest() != hashlib.sha256(after_file.read_bytes()).digest()

    factory = create_session_factory(database_url=os.environ["DATABASE_URL"])

    def snapshot():
        return _snapshot_database(factory)

    def assert_metrics(total, loaded, rejected):
        for label, value in (
            ("전체 상품 행", total), ("정상 변환 예상", loaded), ("Reject 예상", rejected)
        ):
            metric = page.locator('[data-testid="stMetric"]').filter(
                has=page.get_by_text(label, exact=True)
            )
            expect(metric).to_have_count(1)
            expect(metric.locator('[data-testid="stMetricValue"]')).to_have_text(str(value))

    console_errors: list[str] = []
    page_errors: list[str] = []
    page.on("console", lambda message: console_errors.append(message.text)
            if message.type == "error" else None)
    page.on("pageerror", lambda error: page_errors.append(str(error)))
    try:
        page.goto(STREAMLIT_URL, wait_until="domcontentloaded")
        _login_as_operator(page)
        page.get_by_role("tab", name="ETL 적재 이력").click()
        page.get_by_role("combobox", name="ETL 실행 프로필").click()
        page.get_by_role("option", name="마켓플레이스 공급사 샘플", exact=True).click()
        expect(page.locator("body")).to_contain_text(re.compile(r"프로필 버전\s*2"))
        file_input = page.get_by_label("공급사 CSV / XLSX 파일", exact=True).locator('input[type="file"]')
        button = page.get_by_role("button", name="사전 검증", exact=True)
        success = page.get_by_text("사전 검증 완료", exact=True)
        warning = page.get_by_text(
            "일부 상품에 변환 오류가 있습니다. 오류 코드별 건수를 확인한 뒤 원본 파일을 수정해 주세요.",
            exact=True,
        )
        # Scope to the dataframe immediately following the Preflight warning.
        error_table = warning.locator(
            "xpath=ancestor::div[@data-testid='stElementContainer'][1]"
            "/following-sibling::div[@data-testid='stElementContainer'][1]"
        ).locator('[data-testid="stDataFrame"]')
        old_error_cells = page.locator('[data-testid="stDataFrame"]').get_by_role(
            "gridcell", include_hidden=True
        ).filter(has_text=re.compile(r"^(INVALID_PRICE|NEGATIVE_STOCK)$"))
        unsaved_notice = "아직 데이터베이스에 저장하지 않았습니다."
        file_input.set_input_files(str(before_file))
        expect(button).to_be_enabled()
        before = snapshot()
        button.click()
        expect(success).to_be_visible()
        assert_metrics(3, 2, 1)
        expect(warning).to_be_visible()
        expect(page.locator("body")).to_contain_text(unsaved_notice)
        expect(error_table).to_have_count(1)
        error_table.hover()
        # Glide's canvas exposes its displayed cells in a hidden accessibility table.
        expect(error_table.get_by_role("columnheader", include_hidden=True)).to_have_text(
            ["오류 코드", "발생 건수"]
        )
        expect(error_table.get_by_role("gridcell", include_hidden=True)).to_have_text(
            ["INVALID_PRICE", "1", "NEGATIVE_STOCK", "1"]
        )
        after_first = snapshot()
        assert after_first == before

        file_input.set_input_files(str(after_file))
        expect(success).to_have_count(0)
        expect(warning).to_have_count(0)
        expect(error_table).to_have_count(0)
        expect(old_error_cells).to_have_count(0)
        for label in ("전체 상품 행", "정상 변환 예상", "Reject 예상"):
            expect(page.get_by_text(label, exact=True)).to_have_count(0)
        expect(button).to_be_enabled()
        button.click()
        expect(success).to_be_visible()
        assert_metrics(3, 3, 0)
        expect(page.get_by_text("Reject 예상 0건입니다. 변환 오류 코드가 없습니다.", exact=True)).to_be_visible()
        expect(warning).to_have_count(0)
        expect(error_table).to_have_count(0)
        expect(old_error_cells).to_have_count(0)
        expect(page.locator("body")).to_contain_text(unsaved_notice)
        expect(page.get_by_role("button", name="ETL 실행", exact=True)).to_be_enabled()
        after_corrected = snapshot()
        assert before == after_first == after_corrected
        assert source.read_bytes() == original_bytes
        expect(page.locator("body")).not_to_contain_text("StreamlitAPIException")
        assert not console_errors, f"Unexpected browser console errors: {console_errors}"
        assert not page_errors, f"Unexpected browser page errors: {page_errors}"
    except BaseException:
        _preserve_browser_failure_artifacts(page)
        raise
    finally:
        factory.kw["bind"].dispose()


def test_web_etl_preflight_recovery_after_correcting_xlsx_in_real_browser(page, tmp_path):
    from openpyxl import load_workbook
    from playwright.sync_api import expect
    from db.session import create_session_factory

    source = ROOT / "tests" / "fixtures" / "etl" / "sample_marketplace_vendor_mixed.csv"
    original_bytes = source.read_bytes()
    with source.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames
        original_rows = list(reader)
    assert fieldnames is not None
    assert len(original_rows) == 3
    rejected_indices = [
        index for index, row in enumerate(original_rows)
        if row["sku_code"] == "SKU-100-RED-S"
    ]
    assert len(rejected_indices) == 1
    rejected_index = rejected_indices[0]
    assert (
        original_rows[rejected_index]["regular_price"],
        original_rows[rejected_index]["available_qty"],
    ) == ("가격문의", "-1")

    before_file = tmp_path / "before" / "supplier.xlsx"
    after_file = tmp_path / "after" / "supplier.xlsx"
    before_file.parent.mkdir()
    after_file.parent.mkdir()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(fieldnames)
    for row in original_rows:
        worksheet.append([row[column] for column in fieldnames])
    try:
        workbook.save(before_file)
        excel_row = rejected_index + 2
        worksheet.cell(excel_row, fieldnames.index("regular_price") + 1, "59000")
        worksheet.cell(excel_row, fieldnames.index("available_qty") + 1, "3")
        workbook.save(after_file)
    finally:
        workbook.close()

    def read_cells(path):
        saved = load_workbook(path, read_only=True)
        try:
            return list(saved.active.iter_rows(values_only=True))
        finally:
            saved.close()

    before_cells = read_cells(before_file)
    after_cells = read_cells(after_file)
    expected_cells = [tuple(fieldnames)] + [
        tuple(row[column] or None for column in fieldnames) for row in original_rows
    ]
    assert before_cells == expected_cells
    assert len(before_cells) == len(after_cells)
    changed_cells = [
        (row_index, fieldnames[column_index], new_value)
        for row_index, (before_row, after_row) in enumerate(
            zip(before_cells, after_cells, strict=True), start=1
        )
        for column_index, (old_value, new_value) in enumerate(
            zip(before_row, after_row, strict=True)
        )
        if old_value != new_value
    ]
    assert changed_cells == [
        (excel_row, "regular_price", "59000"),
        (excel_row, "available_qty", "3"),
    ]
    assert before_file.name == after_file.name == "supplier.xlsx"
    assert before_file.parent != after_file.parent
    assert hashlib.sha256(before_file.read_bytes()).digest() != hashlib.sha256(
        after_file.read_bytes()
    ).digest()

    factory = create_session_factory(database_url=os.environ["DATABASE_URL"])

    def snapshot():
        return _snapshot_database(factory)

    def assert_metrics(total, loaded, rejected):
        for label, value in (
            ("전체 상품 행", total), ("정상 변환 예상", loaded), ("Reject 예상", rejected)
        ):
            metric = page.locator('[data-testid="stMetric"]').filter(
                has=page.get_by_text(label, exact=True)
            )
            expect(metric).to_have_count(1)
            expect(metric.locator('[data-testid="stMetricValue"]')).to_have_text(str(value))

    console_errors: list[str] = []
    page_errors: list[str] = []
    page.on("console", lambda message: console_errors.append(message.text)
            if message.type == "error" else None)
    page.on("pageerror", lambda error: page_errors.append(str(error)))
    try:
        page.goto(STREAMLIT_URL, wait_until="domcontentloaded")
        _login_as_operator(page)
        page.get_by_role("tab", name="ETL 적재 이력").click()
        page.get_by_role("combobox", name="ETL 실행 프로필").click()
        page.get_by_role("option", name="마켓플레이스 공급사 샘플", exact=True).click()
        expect(page.locator("body")).to_contain_text(re.compile(r"프로필 버전\s*2"))
        file_input = page.get_by_label("공급사 CSV / XLSX 파일", exact=True).locator('input[type="file"]')
        button = page.get_by_role("button", name="사전 검증", exact=True)
        success = page.get_by_text("사전 검증 완료", exact=True)
        warning = page.get_by_text(
            "일부 상품에 변환 오류가 있습니다. 오류 코드별 건수를 확인한 뒤 원본 파일을 수정해 주세요.",
            exact=True,
        )
        # Scope to the dataframe immediately following the Preflight warning.
        error_table = warning.locator(
            "xpath=ancestor::div[@data-testid='stElementContainer'][1]"
            "/following-sibling::div[@data-testid='stElementContainer'][1]"
        ).locator('[data-testid="stDataFrame"]')
        old_error_cells = page.locator('[data-testid="stDataFrame"]').get_by_role(
            "gridcell", include_hidden=True
        ).filter(has_text=re.compile(r"^(INVALID_PRICE|NEGATIVE_STOCK)$"))
        unsaved_notice = "아직 데이터베이스에 저장하지 않았습니다."
        file_input.set_input_files(str(before_file))
        expect(button).to_be_enabled()
        before = snapshot()
        button.click()
        expect(success).to_be_visible()
        assert_metrics(3, 2, 1)
        expect(warning).to_be_visible()
        expect(page.locator("body")).to_contain_text(unsaved_notice)
        expect(error_table).to_have_count(1)
        error_table.hover()
        # Glide's canvas exposes its displayed cells in a hidden accessibility table.
        expect(error_table.get_by_role("columnheader", include_hidden=True)).to_have_text(
            ["오류 코드", "발생 건수"]
        )
        expect(error_table.get_by_role("gridcell", include_hidden=True)).to_have_text(
            ["INVALID_PRICE", "1", "NEGATIVE_STOCK", "1"]
        )
        after_first = snapshot()
        assert after_first == before

        file_input.set_input_files(str(after_file))
        expect(success).to_have_count(0)
        expect(warning).to_have_count(0)
        expect(error_table).to_have_count(0)
        expect(old_error_cells).to_have_count(0)
        for label in ("전체 상품 행", "정상 변환 예상", "Reject 예상"):
            expect(page.get_by_text(label, exact=True)).to_have_count(0)
        expect(button).to_be_enabled()
        button.click()
        expect(success).to_be_visible()
        assert_metrics(3, 3, 0)
        expect(page.get_by_text("Reject 예상 0건입니다. 변환 오류 코드가 없습니다.", exact=True)).to_be_visible()
        expect(warning).to_have_count(0)
        expect(error_table).to_have_count(0)
        expect(old_error_cells).to_have_count(0)
        expect(page.locator("body")).to_contain_text(unsaved_notice)
        expect(page.get_by_role("button", name="ETL 실행", exact=True)).to_be_enabled()
        after_corrected = snapshot()
        assert before == after_first == after_corrected
        assert source.read_bytes() == original_bytes
        expect(page.locator("body")).not_to_contain_text("StreamlitAPIException")
        assert not console_errors, f"Unexpected browser console errors: {console_errors}"
        assert not page_errors, f"Unexpected browser page errors: {page_errors}"
    except BaseException:
        _preserve_browser_failure_artifacts(page)
        raise
    finally:
        factory.kw["bind"].dispose()


def test_xlsx_web_etl_upload_flow_in_real_browser(page, tmp_path):
    fixture_path = tmp_path / "web_etl_upload_vendor.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        [
            "style_id",
            "sku_code",
            "title",
            "category_code",
            "label",
            "regular_price",
            "promo_price",
            "tone",
            "fit_size",
            "available_qty",
            "details",
            "photo",
        ]
    )
    worksheet.append(
        [
            "CG-E2E-WEB-UPLOAD",
            PRODUCT_IDS[0],
            "Synthetic black top",
            "TOP",
            "Synthetic Brand",
            12000,
            10000,
            "BLACK",
            "M",
            3,
            "Synthetic browser XLSX row",
            "black.jpg",
        ]
    )
    worksheet.append(
        [
            "CG-E2E-WEB-UPLOAD",
            PRODUCT_IDS[1],
            "Synthetic white top",
            "TOP",
            "Synthetic Brand",
            15000,
            None,
            "WHITE",
            "L",
            2,
            "Synthetic browser XLSX row",
            "white.jpg",
        ]
    )
    workbook.save(fixture_path)
    workbook.close()

    try:
        _run_web_etl_upload_scenario(
            page,
            fixture_path=fixture_path,
            source_filename=fixture_path.name,
        )
    except BaseException:
        _preserve_browser_failure_artifacts(page)
        raise
