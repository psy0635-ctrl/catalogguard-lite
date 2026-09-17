import io
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from etl import xlsx_reader
from etl.xlsx_reader import XlsxUploadValidationError, read_supplier_xlsx


def _workbook_bytes(rows, *, configure=None) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    if configure is not None:
        configure(workbook, worksheet)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _write_xlsx(tmp_path: Path, content: bytes, name: str = "supplier.xlsx") -> Path:
    path = tmp_path / name
    path.write_bytes(content)
    return path


def _rewrite_zip(content: bytes, *, replacements=None, additions=None) -> bytes:
    replacements = replacements or {}
    additions = additions or {}
    source = io.BytesIO(content)
    target = io.BytesIO()
    with zipfile.ZipFile(source) as input_archive, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as output_archive:
        for member in input_archive.infolist():
            data = replacements.get(member.filename, input_archive.read(member.filename))
            output_archive.writestr(member.filename, data)
        for name, data in additions.items():
            output_archive.writestr(name, data)
    return target.getvalue()


def test_reads_supported_values_and_preserves_source_rows(tmp_path):
    content = _workbook_bytes(
        [
            ["sku", "price", "active", "description"],
            ["00123", 123.0, True, None],
            [None, None, None, None],
            ["SKU-2", 12.5, False, " text "],
        ]
    )

    header, rows, row_numbers, original_bytes = read_supplier_xlsx(
        _write_xlsx(tmp_path, content), ("sku", "price")
    )

    assert header == ["sku", "price", "active", "description"]
    assert rows == [
        {"sku": "00123", "price": "123", "active": "TRUE", "description": ""},
        {"sku": "SKU-2", "price": "12.5", "active": "FALSE", "description": " text "},
    ]
    assert row_numbers == [2, 4]
    assert original_bytes == content


@pytest.mark.parametrize(
    ("rows", "message"),
    [
        ([[None], ["value"]], "헤더"),
        ([[123], ["value"]], "문자열"),
        ([[True], ["value"]], "문자열"),
        ([["sku", " SKU "], ["1", "2"]], "중복"),
        ([["name"], ["value"]], "필수 source"),
        ([["sku"], ["=1+1"]], "수식"),
        ([["sku"], [datetime(2026, 1, 1)]], "날짜 또는 시간"),
    ],
)
def test_rejects_invalid_headers_and_cell_types(tmp_path, rows, message):
    content = _workbook_bytes(rows)
    with pytest.raises(XlsxUploadValidationError, match=message):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_excel_error_cell(tmp_path):
    def configure(_workbook, worksheet):
        worksheet["A2"] = "#DIV/0!"
        worksheet["A2"].data_type = "e"

    content = _workbook_bytes([["sku"], ["placeholder"]], configure=configure)
    with pytest.raises(XlsxUploadValidationError, match="오류 셀"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


@pytest.mark.parametrize(
    "unsupported_value",
    [
        date(2026, 1, 1),
        datetime(2026, 1, 1, 12, 30),
        time(12, 30),
        timedelta(hours=2),
    ],
)
def test_rejects_all_date_and_time_value_types(tmp_path, unsupported_value):
    content = _workbook_bytes([["sku"], [unsupported_value]])
    with pytest.raises(XlsxUploadValidationError, match="날짜 또는 시간"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_multiple_worksheets(tmp_path):
    def configure(workbook, _worksheet):
        workbook.create_sheet("second")

    content = _workbook_bytes([["sku"], ["1"]], configure=configure)
    with pytest.raises(XlsxUploadValidationError, match="정확히 하나"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_rejects_non_visible_only_worksheet(tmp_path, state):
    content = _workbook_bytes([["sku"], ["1"]])
    workbook_xml_name = "xl/workbook.xml"
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        workbook_xml = archive.read(workbook_xml_name).replace(
            b'state="visible"', f'state="{state}"'.encode()
        )
    content = _rewrite_zip(content, replacements={workbook_xml_name: workbook_xml})
    with pytest.raises(XlsxUploadValidationError, match="visible"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_merged_cells(tmp_path):
    def configure(_workbook, worksheet):
        worksheet.merge_cells("A2:B2")

    content = _workbook_bytes([["sku", "name"], ["1", "name"]], configure=configure)
    with pytest.raises(XlsxUploadValidationError, match="병합"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


@pytest.mark.parametrize("content", [b"not-a-zip", b""])
def test_rejects_fake_or_empty_xlsx(tmp_path, content):
    with pytest.raises(XlsxUploadValidationError):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_workbook_without_product_rows(tmp_path):
    content = _workbook_bytes([["sku"]])
    with pytest.raises(XlsxUploadValidationError, match="상품 행"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


@pytest.mark.parametrize("unsafe_name", ["../escape.xml", "/absolute.xml", "C:/drive.xml"])
def test_rejects_unsafe_zip_member_names(tmp_path, unsafe_name):
    content = _rewrite_zip(
        _workbook_bytes([["sku"], ["1"]]),
        additions={unsafe_name: b"synthetic"},
    )
    with pytest.raises(XlsxUploadValidationError, match="구조"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_duplicate_zip_members(tmp_path):
    source = _workbook_bytes([["sku"], ["1"]])
    duplicate = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as input_archive, zipfile.ZipFile(
        duplicate, "w", compression=zipfile.ZIP_DEFLATED
    ) as output_archive:
        for member in input_archive.infolist():
            output_archive.writestr(member.filename, input_archive.read(member.filename))
        with pytest.warns(UserWarning, match="Duplicate name"):
            output_archive.writestr("xl/workbook.xml", b"duplicate")
    with pytest.raises(XlsxUploadValidationError, match="구조"):
        read_supplier_xlsx(_write_xlsx(tmp_path, duplicate.getvalue()), ("sku",))


def test_rejects_unsupported_zip_compression(tmp_path):
    source = _workbook_bytes([["sku"], ["1"]])
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as input_archive, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_BZIP2
    ) as output_archive:
        for member in input_archive.infolist():
            output_archive.writestr(member.filename, input_archive.read(member.filename))
    with pytest.raises(XlsxUploadValidationError, match="압축 형식"):
        read_supplier_xlsx(_write_xlsx(tmp_path, target.getvalue()), ("sku",))


def test_rejects_encrypted_zip_member(tmp_path):
    content = bytearray(_workbook_bytes([["sku"], ["1"]]))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        member = archive.getinfo("[Content_Types].xml")
        local_flag_offset = member.header_offset + 6
    filename = b"[Content_Types].xml"
    central_name_offset = content.rfind(filename)
    central_header_offset = central_name_offset - 46
    assert content[central_header_offset : central_header_offset + 4] == b"PK\x01\x02"
    central_flag_offset = central_header_offset + 8
    for offset in (local_flag_offset, central_flag_offset):
        flag_bits = int.from_bytes(content[offset : offset + 2], "little") | 0x1
        content[offset : offset + 2] = flag_bits.to_bytes(2, "little")

    with pytest.raises(XlsxUploadValidationError, match="암호화"):
        read_supplier_xlsx(_write_xlsx(tmp_path, bytes(content)), ("sku",))


def test_corrupt_required_member_is_reported_as_safe_validation_error(tmp_path):
    content = bytearray(_workbook_bytes([["sku"], ["1"]]))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        member = archive.getinfo("[Content_Types].xml")
        # Corrupt compressed payload bytes without changing the central directory metadata.
        offset = member.header_offset + len(member.FileHeader())
    content[offset] ^= 0xFF
    with pytest.raises(XlsxUploadValidationError, match="구조"):
        read_supplier_xlsx(_write_xlsx(tmp_path, bytes(content)), ("sku",))


def test_rejects_macro_payload_even_with_xlsx_extension(tmp_path):
    content = _rewrite_zip(
        _workbook_bytes([["sku"], ["1"]]),
        additions={"xl/vbaProject.bin": b"synthetic-test-payload"},
    )
    with pytest.raises(XlsxUploadValidationError, match="매크로"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_external_workbook_link(tmp_path):
    content = _rewrite_zip(
        _workbook_bytes([["sku"], ["1"]]),
        additions={"xl/externalLinks/externalLink1.xml": b"<externalLink />"},
    )
    with pytest.raises(XlsxUploadValidationError, match="외부 workbook"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_cells_beyond_column_limit(tmp_path):
    def configure(_workbook, worksheet):
        worksheet.cell(row=2, column=xlsx_reader.MAX_XLSX_COLUMNS + 1, value="too-wide")

    content = _workbook_bytes([["sku"], ["1"]], configure=configure)
    with pytest.raises(XlsxUploadValidationError, match="열 개수"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_sparse_row_beyond_iteration_budget(tmp_path):
    def configure(_workbook, worksheet):
        worksheet.cell(row=xlsx_reader.MAX_CSV_ROWS + 2, column=1, value="too-far")

    content = _workbook_bytes([["sku"], ["1"]], configure=configure)
    with pytest.raises(XlsxUploadValidationError, match="행 위치"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_rejects_zip_entry_and_size_limits(tmp_path, monkeypatch):
    content = _workbook_bytes([["sku"], ["1"]])
    monkeypatch.setattr(xlsx_reader, "MAX_XLSX_ZIP_ENTRIES", 1)
    with pytest.raises(XlsxUploadValidationError, match="구조가 너무 큽니다"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_applies_existing_compressed_upload_limit_before_opening_workbook(tmp_path):
    content = b"x" * (xlsx_reader.MAX_UPLOAD_SIZE_BYTES + 1)
    with pytest.raises(XlsxUploadValidationError, match="파일 크기"):
        read_supplier_xlsx(_write_xlsx(tmp_path, content), ("sku",))


def test_accepts_ten_thousand_products_and_rejects_next_row(tmp_path):
    rows = [["sku"], *[[f"SKU-{index}"] for index in range(xlsx_reader.MAX_CSV_ROWS)]]
    content = _workbook_bytes(rows)
    _header, products, row_numbers, _bytes = read_supplier_xlsx(
        _write_xlsx(tmp_path, content, "boundary.xlsx"), ("sku",)
    )
    assert len(products) == xlsx_reader.MAX_CSV_ROWS
    assert row_numbers[-1] == xlsx_reader.MAX_CSV_ROWS + 1

    oversized_content = _workbook_bytes([*rows, ["one-too-many"]])
    with pytest.raises(XlsxUploadValidationError, match="행"):
        read_supplier_xlsx(
            _write_xlsx(tmp_path, oversized_content, "oversized.xlsx"), ("sku",)
        )


def test_rejects_single_member_total_and_compression_ratio_limits(tmp_path, monkeypatch):
    content = _workbook_bytes([["sku"], ["1"]])
    path = _write_xlsx(tmp_path, content)

    monkeypatch.setattr(xlsx_reader, "MAX_XLSX_ZIP_MEMBER_BYTES", 1)
    with pytest.raises(XlsxUploadValidationError, match="내부 파일"):
        read_supplier_xlsx(path, ("sku",))

    monkeypatch.setattr(xlsx_reader, "MAX_XLSX_ZIP_MEMBER_BYTES", 20 * 1024 * 1024)
    monkeypatch.setattr(xlsx_reader, "MAX_XLSX_ZIP_TOTAL_BYTES", 1)
    with pytest.raises(XlsxUploadValidationError, match="압축 해제"):
        read_supplier_xlsx(path, ("sku",))

    monkeypatch.setattr(xlsx_reader, "MAX_XLSX_ZIP_TOTAL_BYTES", 50 * 1024 * 1024)
    monkeypatch.setattr(xlsx_reader, "MAX_XLSX_COMPRESSION_RATIO", 1)
    with pytest.raises(XlsxUploadValidationError, match="압축률"):
        read_supplier_xlsx(path, ("sku",))
