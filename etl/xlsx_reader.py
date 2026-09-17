import io
import math
import re
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path, PurePosixPath

from defusedxml import ElementTree as SafeElementTree
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from config.settings import (
    MAX_CSV_ROWS,
    MAX_UPLOAD_SIZE_BYTES,
    MAX_XLSX_COLUMNS,
    MAX_XLSX_COMPRESSION_RATIO,
    MAX_XLSX_ZIP_ENTRIES,
    MAX_XLSX_ZIP_MEMBER_BYTES,
    MAX_XLSX_ZIP_TOTAL_BYTES,
)
from core.upload_validator import find_duplicate_columns


class XlsxUploadValidationError(ValueError):
    """Safe, user-facing validation failure for a supplier XLSX."""


_REQUIRED_OOXML_MEMBERS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
}
_ALLOWED_COMPRESSION_TYPES = {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
_CELL_REFERENCE_RE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
_SPREADSHEET_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _invalid(message: str) -> XlsxUploadValidationError:
    return XlsxUploadValidationError(message)


def validate_xlsx_filename(filename: str | None) -> None:
    if not filename:
        raise _invalid("XLSX 파일만 업로드할 수 있습니다.")
    basename = str(filename).replace("\\", "/").split("/")[-1].strip()
    if (
        "\x00" in basename
        or not basename.casefold().endswith(".xlsx")
        or not basename[:-5].strip()
    ):
        raise _invalid("XLSX 파일만 업로드할 수 있습니다.")


def _validate_member_name(name: str) -> None:
    normalized = name.replace("\\", "/")
    path = PurePosixPath(normalized)
    if (
        not normalized
        or normalized.startswith("/")
        or re.match(r"^[A-Za-z]:", normalized)
        or ".." in path.parts
    ):
        raise _invalid("XLSX 파일 구조가 올바르지 않습니다.")


def _validate_zip_archive(file_bytes: bytes) -> tuple[zipfile.ZipFile, set[str]]:
    if not file_bytes:
        raise _invalid("업로드한 파일이 비어 있습니다.")
    if len(file_bytes) > MAX_UPLOAD_SIZE_BYTES:
        raise _invalid("파일 크기가 너무 큽니다. 최대 5MB까지 업로드할 수 있습니다.")

    try:
        archive = zipfile.ZipFile(io.BytesIO(file_bytes))
        members = archive.infolist()
    except (zipfile.BadZipFile, OSError, ValueError) as error:
        raise _invalid("올바른 XLSX 파일이 아닙니다.") from error

    try:
        if len(members) > MAX_XLSX_ZIP_ENTRIES:
            raise _invalid("XLSX 파일 구조가 너무 큽니다.")
        names: set[str] = set()
        total_size = 0
        for member in members:
            _validate_member_name(member.filename)
            if member.filename in names:
                raise _invalid("XLSX 파일 구조가 올바르지 않습니다.")
            names.add(member.filename)
            if member.flag_bits & 0x1:
                raise _invalid("암호화된 XLSX 파일은 지원하지 않습니다.")
            if member.compress_type not in _ALLOWED_COMPRESSION_TYPES:
                raise _invalid("XLSX 압축 형식을 지원하지 않습니다.")
            if member.file_size > MAX_XLSX_ZIP_MEMBER_BYTES:
                raise _invalid("XLSX 내부 파일이 너무 큽니다.")
            total_size += member.file_size
            if total_size > MAX_XLSX_ZIP_TOTAL_BYTES:
                raise _invalid("XLSX 압축 해제 크기가 너무 큽니다.")
            if member.file_size and (
                member.compress_size == 0
                or member.file_size / member.compress_size > MAX_XLSX_COMPRESSION_RATIO
            ):
                raise _invalid("XLSX 압축률이 허용 범위를 초과했습니다.")

        if not _REQUIRED_OOXML_MEMBERS.issubset(names):
            raise _invalid("XLSX OOXML 구조가 올바르지 않습니다.")
        if not any(name.startswith("xl/worksheets/") and name.endswith(".xml") for name in names):
            raise _invalid("XLSX worksheet를 찾을 수 없습니다.")
        return archive, names
    except Exception:
        archive.close()
        raise


def _parse_xml_member(archive: zipfile.ZipFile, name: str):
    try:
        with archive.open(name) as source:
            return SafeElementTree.parse(source).getroot()
    except Exception as error:
        raise _invalid("XLSX XML 구조가 올바르지 않습니다.") from error


def _read_member_bytes(archive: zipfile.ZipFile, name: str) -> bytes:
    try:
        return archive.read(name)
    except Exception as error:
        raise _invalid("XLSX 파일 구조가 올바르지 않습니다.") from error


def _validate_package_relationships(archive: zipfile.ZipFile, names: set[str]) -> None:
    lowered_names = {name.casefold() for name in names}
    if any("vbaproject" in name for name in lowered_names):
        raise _invalid("매크로가 포함된 XLSX 파일은 지원하지 않습니다.")
    if any(name.startswith("xl/externallinks/") for name in lowered_names):
        raise _invalid("외부 workbook link가 포함된 XLSX 파일은 지원하지 않습니다.")

    content_types = _read_member_bytes(archive, "[Content_Types].xml").lower()
    if b"macroenabled" in content_types or b"vbaproject" in content_types:
        raise _invalid("매크로가 포함된 XLSX 파일은 지원하지 않습니다.")

    for name in names:
        if not name.endswith(".rels"):
            continue
        root = _parse_xml_member(archive, name)
        for relationship in root.findall(f"{{{_RELATIONSHIP_NAMESPACE}}}Relationship"):
            relationship_type = relationship.attrib.get("Type", "").casefold()
            if "vbaproject" in relationship_type:
                raise _invalid("매크로가 포함된 XLSX 파일은 지원하지 않습니다.")
            if "externallink" in relationship_type:
                raise _invalid("외부 workbook link가 포함된 XLSX 파일은 지원하지 않습니다.")


def _validate_worksheet_xml(archive: zipfile.ZipFile, names: set[str]) -> None:
    worksheet_names = sorted(
        name for name in names if name.startswith("xl/worksheets/") and name.endswith(".xml")
    )
    for name in worksheet_names:
        root = _parse_xml_member(archive, name)
        if root.find(f".//{{{_SPREADSHEET_NAMESPACE}}}mergeCell") is not None:
            raise _invalid("XLSX 파일에 병합된 셀이 포함되어 있습니다.")
        if root.find(f".//{{{_SPREADSHEET_NAMESPACE}}}f") is not None:
            raise _invalid("XLSX 파일에 수식 셀이 포함되어 있습니다.")
        cell_count = 0
        for cell in root.iter(f"{{{_SPREADSHEET_NAMESPACE}}}c"):
            reference = cell.attrib.get("r", "")
            match = _CELL_REFERENCE_RE.fullmatch(reference)
            if match is None:
                raise _invalid("XLSX 셀 위치가 올바르지 않습니다.")
            column_index = column_index_from_string(match.group(1))
            row_index = int(match.group(2))
            if column_index > MAX_XLSX_COLUMNS:
                raise _invalid("XLSX 열 개수가 허용 범위를 초과했습니다.")
            # read_only 모드가 거대한 sparse 좌표 사이를 채우지 않도록 실제 좌표도 제한합니다.
            if row_index > MAX_CSV_ROWS + 1:
                raise _invalid("XLSX 행 위치가 허용 범위를 초과했습니다.")
            cell_count += 1
            if cell_count > (MAX_CSV_ROWS + 1) * MAX_XLSX_COLUMNS:
                raise _invalid("XLSX 셀 개수가 허용 범위를 초과했습니다.")


def _stringify_cell(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "f":
        raise _invalid("XLSX 파일에 수식 셀이 포함되어 있습니다.")
    if cell.data_type == "e":
        raise _invalid("XLSX 파일에 오류 셀이 포함되어 있습니다.")
    if cell.is_date or isinstance(value, (date, datetime, time, timedelta)):
        raise _invalid("XLSX 파일에 날짜 또는 시간 셀이 포함되어 있습니다.")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise _invalid("XLSX 파일에 유효하지 않은 숫자가 포함되어 있습니다.")
        decimal_value = Decimal(str(value))
        if decimal_value == decimal_value.to_integral_value():
            return format(decimal_value.quantize(Decimal(1)), "f")
        return format(decimal_value.normalize(), "f")
    raise _invalid("XLSX 파일에 지원하지 않는 셀 형식이 포함되어 있습니다.")


def read_supplier_xlsx(
    input_path: Path,
    required_columns: tuple[str, ...],
) -> tuple[list[str], list[dict[str, str]], list[int], bytes]:
    try:
        file_bytes = input_path.read_bytes()
    except FileNotFoundError as error:
        raise _invalid("Input XLSX file was not found") from error
    except OSError as error:
        raise _invalid("Input XLSX file could not be read") from error

    validate_xlsx_filename(input_path.name)
    archive, names = _validate_zip_archive(file_bytes)
    try:
        _validate_package_relationships(archive, names)
        _validate_worksheet_xml(archive, names)
    finally:
        archive.close()

    try:
        workbook = load_workbook(
            input_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise _invalid("올바른 XLSX workbook이 아닙니다.") from error

    try:
        if len(workbook.worksheets) != 1:
            raise _invalid("XLSX 파일에는 worksheet가 정확히 하나 있어야 합니다.")
        worksheet = workbook.worksheets[0]
        if worksheet.sheet_state != "visible":
            raise _invalid("XLSX worksheet는 visible 상태여야 합니다.")

        iterator = worksheet.iter_rows(max_col=MAX_XLSX_COLUMNS)
        try:
            header_cells = next(iterator)
        except StopIteration as error:
            raise _invalid("XLSX 파일이 비어 있습니다.") from error

        last_header_index = next(
            (index for index in range(len(header_cells) - 1, -1, -1) if header_cells[index].value is not None),
            -1,
        )
        if last_header_index < 0:
            raise _invalid("XLSX 헤더가 비어 있습니다.")
        header: list[str] = []
        for cell in header_cells[: last_header_index + 1]:
            if not isinstance(cell.value, str) or cell.data_type == "f":
                raise _invalid("XLSX 헤더는 문자열이어야 합니다.")
            header.append(cell.value.strip())
        if not header or any(not column for column in header):
            raise _invalid("XLSX 헤더에 빈 열이 있습니다.")
        if find_duplicate_columns(header):
            raise _invalid("XLSX 헤더에 중복된 열이 있습니다.")
        if any(column not in header for column in required_columns):
            raise _invalid("XLSX 파일에 필수 source 열이 없습니다.")

        rows: list[dict[str, str]] = []
        row_numbers: list[int] = []
        for row_cells in iterator:
            all_values = [_stringify_cell(cell) for cell in row_cells]
            if any(value.strip() for value in all_values[len(header) :]):
                raise _invalid("XLSX 상품 행의 열 개수가 헤더와 일치하지 않습니다.")
            values = all_values[: len(header)]
            if all(not value.strip() for value in values):
                continue
            rows.append(dict(zip(header, values, strict=True)))
            row_numbers.append(row_cells[0].row)
            if len(rows) > MAX_CSV_ROWS:
                raise _invalid("XLSX 상품 행이 허용 개수를 초과했습니다.")
        if not rows:
            raise _invalid("XLSX 파일에 상품 행이 없습니다.")
        return header, rows, row_numbers, file_bytes
    finally:
        workbook.close()
